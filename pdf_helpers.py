import os
import time
import tabula
import camelot
import pdfplumber
import pandas as pd
from datetime import datetime
from openpyxl import Workbook, load_workbook


def format_number(num):
    if num >= 1_000_000_000_000:
        return f'{num / 1_000_000_000_000:.1f}T'
    if num >= 1_000_000_000:
        return f'{num / 1_000_000_000:.1f}B'
    elif num >= 1_000_000:
        return f'{num / 1_000_000:.1f}M'
    elif num >= 1_000:
        return f'{num / 1_000:.1f}K'
    else:
        return str(num)


def convert_seconds(seconds):
    hours = seconds // 3600
    minutes = (seconds % 3600) // 60
    seconds = seconds % 60
    return f"{int(hours):02}:{int(minutes):02}:{round(seconds, 2):02}"


def process_pdf_plumber(args, args2):
    pdf, page, total_pages, prefix, directory = args
    total_files, sum_pages, current_pages = args2

    output = f"{directory}/data_{prefix}{str(page).zfill(7)}.xlsx"
    st = time.time()
    if (not os.path.exists(output)):
        content = pdf.pages[page - 1]
        table = content.extract_table()
        df = pd.DataFrame(table[1:], columns=table[0])
        df.to_excel(output, index=False)

    et = time.time() - st
    estimate = convert_seconds((sum_pages - current_pages) * et)
    file_status = "{:<8}".format(f"[{prefix + 1}/{total_files}]")
    page_status = "{:<12}".format(f"{page}/{total_pages - 1}")
    et_status = "{:<6}".format(f"{round(et, 2)}")
    estimate_status = "{:<22}".format(f"{estimate}")
    process = round(current_pages / sum_pages * 100, 2)
    process_status = "{:<6}".format(f"{process}")
    print(f"{file_status} Export: {page_status} | Time: {
          et_status} | Process: {process_status} | Estimate: {estimate_status}")


def process_pdf_tabula(args, args2):
    input, page, total_pages, prefix, directory = args
    total_files, sum_pages, current_pages = args2

    output = f"{directory}/data_{prefix}{str(page).zfill(7)}.xlsx"
    st = time.time()
    if (not os.path.exists(output)):
        df = tabula.read_pdf(input, pages=page, stream=True)[0]
        df.to_excel(output, index=False)

    et = time.time() - st
    estimate = convert_seconds((sum_pages - current_pages) * et)
    file_status = "{:<8}".format(f"[{prefix + 1}/{total_files}]")
    page_status = "{:<12}".format(f"{page}/{total_pages - 1}")
    et_status = "{:<6}".format(f"{round(et, 2)}")
    estimate_status = "{:<22}".format(f"{estimate}")
    process = round(current_pages / sum_pages * 100, 2)
    process_status = "{:<6}".format(f"{process}")
    print(f"{file_status} Export: {page_status} | Time: {
          et_status} | Process: {process_status} | Estimate: {estimate_status}")


def process_pdf_camelot(args):
    input, page, total_pages, prefix, directory = args
    output = f"{directory}/data_{prefix}{str(page).zfill(7)}.csv"
    st = time.time()
    if (not os.path.exists(output)):
        cm = camelot.read_pdf(input, pages=str(page), flavor='stream')[0]
        cm.to_csv(output)
    et = time.time() - st
    estimate = (total_pages - page) * et / 60 / 60
    print(f"Export: {page}/{total_pages - 1} | {et} | Estimate: {estimate}")


def process_excel_plumber(df, data, total_rows):
    print("process_excel_plumber")


def standardize_message(message: str):
    message = message.replace('None', '')
    message = ' '.join(message.split())
    return message


date_format = "%d/%m/%Y %H:%M:%S"
tabulaDatetimeIndex = 1
tabulaAmountIndex = 3
tabulaMessageIndex = 5


def standardize_excel_tabula(input, output, override=False):
    try:
        if (not override and os.path.exists(output)):
            print(f"{output}", f"{input}")
            return
        workbook = load_workbook(input)
        sheet = workbook.active
        max_row = sheet.max_row
        data = list()
        for row in range(1, max_row + 1):
            amount = str(sheet.cell(
                row=row, column=tabulaAmountIndex).value).replace(".", "")
            print(amount)
            if (amount.isnumeric()):
                issue_date = datetime.strptime(str(sheet.cell(
                    row=row - 1, column=tabulaDatetimeIndex).value) + " 00:00:00", date_format).strftime('%Y-%m-%d %H:%M:%S')
                message = str(sheet.cell(
                    row=row, column=tabulaMessageIndex).value)
                data.append([issue_date, int(amount),
                            standardize_message(message)])
            elif (len(data) > 0):
                message = str(sheet.cell(
                    row=row, column=tabulaMessageIndex).value)
                if (message != "nan"):
                    data[len(data) - 1][2] = standardize_message(data[len(data) - 1][2] +
                                                                 " " + message)
        newWorkbook = Workbook()
        newSheet = newWorkbook.active
        for row in data:
            newSheet.append(row)
        newWorkbook.save(output)
        print(f"{output}", f"{input}")
    except Exception as e:
        print(f"ERROR: {output}: {str(e)}")


tabula2DatetimeIndex = 2
tabula2AmountIndex = 4
tabula2MessageIndex = 5


def standardize_excel_tabula2(input, output, override=False):
    try:
        if (not override and os.path.exists(output)):
            print(f"{output}", f"{input}")
            return
        workbook = load_workbook(input)
        sheet = workbook.active
        max_row = sheet.max_row
        data = list()
        for row in range(1, max_row + 1):
            amount = str(sheet.cell(
                row=row, column=tabula2AmountIndex).value).replace(".", "").split(',')[0]
            if (amount.isnumeric()):
                issue_date = str(sheet.cell(
                    row=row, column=tabula2DatetimeIndex).value)
                if (len(issue_date) <= 10):
                    issue_date = issue_date + " 00:00:00"
                elif (len(issue_date) <= 17):
                    issue_date = issue_date[0:10] + " 0" + issue_date[10:]
                issue_date = datetime.strptime(
                    issue_date, date_format).strftime('%Y-%m-%d %H:%M:%S')
                message = str(sheet.cell(
                    row=row, column=tabula2MessageIndex).value).replace("None", "")
                message = message + " " + \
                    str(sheet.cell(row=row, column=tabula2MessageIndex +
                        1).value).replace("None", "")
                data.append([issue_date, int(amount),
                            standardize_message(message)])
            elif (len(data) > 0):
                message = str(sheet.cell(
                    row=row, column=tabula2MessageIndex + 1).value).replace("None", "")
                if (message != "nan"):
                    data[len(data) - 1][2] = standardize_message(data[len(data) - 1][2] +
                                                                 " " + message)
        newWorkbook = Workbook()
        newSheet = newWorkbook.active
        for row in data:
            newSheet.append(row)
        newWorkbook.save(output)
        print(f"{output}", f"{input}")
    except Exception as e:
        print(f"ERROR: {output}: {str(e)}")


plumberDatetimeIndex = 2
plumberAmountIndex = 3
plumberMessageIndex = 4


def standardize_excel_plumber(input, output, override=False):
    try:
        if (not override and os.path.exists(output)):
            print(f"{output}", f"{input}")
            return
        workbook = load_workbook(input)
        sheet = workbook.active
        max_row = sheet.max_row
        newWorkbook = Workbook()
        newSheet = newWorkbook.active
        for row in range(1, max_row + 1):
            amount = str(sheet.cell(row=row, column=plumberAmountIndex).value).replace(
                ".", "").split(',')[0]
            if (amount.isnumeric()):
                issue_date = str(sheet.cell(
                    row=row, column=plumberDatetimeIndex).value)
                if (len(issue_date) <= 10):
                    issue_date = issue_date + " 00:00:00"
                issue_date = datetime.strptime(
                    issue_date, date_format).strftime('%Y-%m-%d %H:%M:%S')
                message = str(sheet.cell(row=row, column=plumberMessageIndex).value).replace('\n', ' ').replace(
                    "\"", "") + str(sheet.cell(row=row, column=plumberMessageIndex + 1).value).replace('\n', ' ').replace("None", "")
                newSheet.append(
                    [issue_date, int(amount), standardize_message(message)])
        newWorkbook.save(output)
        print(f"{output}", f"{input}")
    except Exception as e:
        print(f"ERROR: {output}: {str(e)}")


def combine_tokens(entities):
    combined_entities = []
    current_entity = ""
    current_label = None

    for entity in entities:
        word = entity['word']
        label = entity['entity']

        if word.startswith("##"):
            current_entity += word[2:]
        else:
            if current_entity:
                combined_entities.append((current_entity, current_label))
            current_entity = word
            current_label = label

    if current_entity:
        combined_entities.append((current_entity, current_label))

    return combined_entities


def filter_per(item):
    word, label = item
    if label == 'B-PER':
        return True
    if label == 'I-PER':
        return True
    return False


def merge_information(entities):
    full_names = []
    current_name = ''
    filtered_array = list(filter(filter_per, entities))
    if len(filtered_array) == 0:
        return []

    step = 0
    for word, tag in filtered_array:
        if tag == 'B-PER':
            if step == 0:
                step = 1
            elif step == 2:
                step = 1
                full_names.append(" ".join(current_name.split()))
                current_name = ''
            current_name = current_name + ' ' + word
        elif tag == 'I-PER':
            if step == 1:
                step = 2
            if step == 2:
                current_name = current_name + ' ' + word
    if len(current_name) > 0:
        full_names.append(" ".join(current_name.split()))
    return full_names


def extract_information(entities):
    first_names = list()
    last_names = list()
    for entity in entities:
        word = entity[0]
        label = entity[1]
        if (label == 'B-PER'):
            first_names.append(word)
        elif (label == 'I-PER'):
            last_names.append(word)
    return [first_names, last_names]


def ai_extract_information(text, ner_pipeline):
    text = str(text).lower()
    entities = ner_pipeline(text)
    entities = combine_tokens(entities)
    return extract_information(entities)


def ai_extract_fullname(text, ner_pipeline):
    text = str(text).lower()
    entities = ner_pipeline(text)
    entities = combine_tokens(entities)
    return merge_information(entities)
